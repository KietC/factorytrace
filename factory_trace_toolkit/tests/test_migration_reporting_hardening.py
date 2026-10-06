"""English: Regress lexical path hardening, legacy logs, and edited report linting.

中文：回归词法路径加固、历史日志兼容与编辑后报告检查；临时案件验证迁移不能越界，旧概率语义不能回流。
"""
from __future__ import annotations

import os
import subprocess
import tempfile
import unittest
from pathlib import Path

from factorytrace.audit import audit_case
from factorytrace.init_case import create_case
from factorytrace.reporting import lint_report_file, lint_report_text
from factorytrace.v2_commands import lint_case_reports, migrate_case_directory


class MigrationPathHardeningTests(unittest.TestCase):
    def _make_directory_link(self, link: Path, target: Path) -> None:
        if os.name == "nt":
            result = subprocess.run(
                ["cmd.exe", "/d", "/c", "mklink", "/J", str(link), str(target)],
                capture_output=True,
                text=True,
                check=False,
            )
            if result.returncode != 0:
                self.skipTest(f"could not create Windows junction: {result.stderr}")
        else:
            os.symlink(target, link, target_is_directory=True)

    def test_migration_rejects_source_root_reparse_before_resolve(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            real_source = create_case("real-source", root)
            linked_source = root / "linked-source"
            self._make_directory_link(linked_source, real_source)

            with self.assertRaisesRegex(ValueError, "reparse|junction|symbolic"):
                migrate_case_directory(
                    linked_source, destination_root=root / "migrated"
                )
            self.assertFalse((root / "migrated").exists())

    def test_migration_rejects_destination_parent_reparse_before_resolve(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = create_case("source", root)
            real_parent = root / "real-destination-parent"
            real_parent.mkdir()
            linked_parent = root / "linked-destination-parent"
            self._make_directory_link(linked_parent, real_parent)

            with self.assertRaisesRegex(ValueError, "reparse|junction|symbolic"):
                migrate_case_directory(
                    source, destination_root=linked_parent / "migrated"
                )
            self.assertFalse((real_parent / "migrated").exists())

    def test_canonical_copy_excludes_runtime_and_release_segments(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = create_case("source", root)
            (source / "evidence" / "valid.pdf").write_bytes(b"frozen evidence")
            (source / "logs" / "audit.json").write_text("{}", encoding="utf-8")
            excluded = [
                source / "evidence" / ".venv" / "leak.txt",
                source / "evidence" / "node_modules" / "leak.txt",
                source / "evidence" / "cache" / "leak.txt",
                source / "evidence" / "release" / "leak.txt",
                source / "evidence" / "renders" / "leak.txt",
                source / "evidence" / "staging" / "leak.txt",
                source / "logs" / "release.zip",
            ]
            for path in excluded:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(b"must not migrate")

            destination = root / "migrated"
            receipt = migrate_case_directory(source, destination_root=destination)

            self.assertTrue((destination / "notes.md").is_file())
            self.assertTrue((destination / "evidence" / "valid.pdf").is_file())
            self.assertTrue((destination / "logs" / "audit.json").is_file())
            for source_path in excluded:
                relative = source_path.relative_to(source)
                self.assertFalse((destination / relative).exists(), str(relative))
            skipped = {row["path"] for row in receipt["skipped_entries"]}
            self.assertIn("logs/release.zip", skipped)


class AuditCompatibilityTests(unittest.TestCase):
    def test_audit_accepts_legacy_logs_commands_log_with_warning(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case = create_case("legacy-command-log", Path(temporary))
            legacy_path = case / "logs" / "commands.log"
            (case / "commands.log").replace(legacy_path)

            report = audit_case(case, "structure")

            self.assertFalse(
                any("missing file: commands.log" in item for item in report["errors"])
            )
            self.assertTrue(
                any("logs/commands.log" in item for item in report["warnings"])
            )


class LegacyProbabilityReportTests(unittest.TestCase):
    def test_lint_rejects_legacy_multi_axis_probability_semantics(self) -> None:
        examples = [
            "四轴百分比：A 85% / B 65% / C 30% / D 0%",
            "四维概率定义",
            "最高综合采购适配估计 0.68",
            "候选厂家｜综合适配与四维角色置信",
            "候选甲 | 85% | 65% | 30% | 0%",
        ]
        for example in examples:
            with self.subTest(example=example):
                self.assertTrue(lint_report_text(example))

    def test_lint_rejects_old_main_report_patterns_in_md_xlsx_docx(self) -> None:
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("reports extra is not installed")

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            markdown = root / "old-main.md"
            markdown.write_text(
                "# 主报告\n\n## 四维概率定义\n\n"
                "| 候选 | A | B | C | D |\n|---|---:|---:|---:|---:|\n"
                "| 候选甲 | 85% | 65% | 30% | 0% |\n",
                encoding="utf-8",
            )

            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["候选厂家｜综合适配与四维角色置信"])
            sheet.append(
                [
                    "候选厂家",
                    "综合采购适配估计",
                    "A 谱系同源",
                    "B 制造",
                    "C 模具",
                    "D 双证",
                ]
            )
            sheet.append(["候选甲", 0.68, 0.85, 0.65, 0.30, 0.0])
            xlsx = root / "old-main.xlsx"
            workbook.save(xlsx)
            workbook.close()

            document = Document()
            document.add_heading("四维概率定义", level=1)
            table = document.add_table(rows=2, cols=5)
            for index, value in enumerate(["候选", "A", "B", "C", "D"]):
                table.rows[0].cells[index].text = value
            for index, value in enumerate(["候选甲", "85%", "65%", "30%", "0%"]):
                table.rows[1].cells[index].text = value
            docx = root / "old-main.docx"
            document.save(docx)

            self.assertTrue(lint_report_file(markdown))
            self.assertTrue(lint_report_file(xlsx))
            self.assertTrue(lint_report_file(docx))

    def test_default_case_lint_discovers_reports_and_rankings_recursively(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            case = create_case("lint-discovery", Path(temporary))
            report_dir = case / "output" / "reports"
            ranking_dir = case / "output" / "rankings"
            report_dir.mkdir(parents=True, exist_ok=True)
            ranking_dir.mkdir(parents=True, exist_ok=True)
            report = report_dir / "old-main.md"
            ranking = ranking_dir / "old-ranking.md"
            report.write_text("## 四维概率定义\n", encoding="utf-8")
            ranking.write_text("最高综合采购适配估计 0.68\n", encoding="utf-8")

            result = lint_case_reports(case)

            self.assertEqual(result["status"], "FAIL")
            self.assertIn(str(report), result["checked"])
            self.assertIn(str(ranking), result["checked"])


if __name__ == "__main__":
    unittest.main()
