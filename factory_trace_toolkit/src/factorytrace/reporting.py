"""English: Render and lint consistent Markdown, spreadsheet, and Word evidence reports.

中文：生成并检查 Markdown、Excel 与 Word 证据报告；编辑后的 Office 文件也须复检，链接与排序指标不能掩盖缺口。
"""

from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable

from .common import (
    atomic_write_json,
    atomic_write_text,
    sanitize_url_for_record,
    sha256_file,
    utc_now,
    write_csv,
)
from .scoring import score_candidate


TRACE_STAGE_ORDER = {
    "S0_UNLINKED": 0,
    "S1_EXACT_PRODUCT_MATCH": 1,
    "S2_AUTHORIZED_SITE": 2,
    "S3_PROCESS_CONFIRMED": 3,
    "S4_BATCH_LINKED": 4,
}

REPORT_COLUMNS = [
    "candidate_id",
    "candidate",
    "verified_mainland_site",
    "target_role",
    "trace_stage",
    "analytic_confidence",
    "evidence_sufficiency_score",
    "capability_fit_score",
    "ul_state",
    "watermark_state",
    "verification_priority",
    "procurement_utility",
    "critical_gap",
    "contact",
    "contact_status",
    "website_or_product_page",
    "contact_source_url",
    "contact_evidence_path",
    "last_verified_utc",
    "evidence_source_urls",
]

REPORT_LABELS = {
    "candidate_id": "候选 ID",
    "candidate": "候选",
    "verified_mainland_site": "大陆场址",
    "target_role": "目标角色",
    "trace_stage": "溯源阶段",
    "analytic_confidence": "分析置信度",
    "evidence_sufficiency_score": "证据闭合度",
    "capability_fit_score": "能力适配",
    "ul_state": "UL",
    "watermark_state": "WaterMark",
    "verification_priority": "核验优先级",
    "procurement_utility": "采购效用",
    "critical_gap": "关键缺口",
    "contact": "联系方式",
    "contact_status": "联系方式状态",
    "website_or_product_page": "已核验官网/产品页",
    "contact_source_url": "联系方式来源",
    "contact_evidence_path": "联系方式冻结证据",
    "last_verified_utc": "最后核验时间 UTC",
    "evidence_source_urls": "其他证据来源 URL",
}

URL_RE = re.compile(r"https?://[^\s<>\[\](){}\"'，。；、]+", re.IGNORECASE)

FORBIDDEN_PERCENT_RE = re.compile(
    r"(?i)(?:厂家|制造商|源头|source\s*factory|manufacturer)"
    r".{0,32}(?:概率|可能性|置信度|probability|likelihood|confidence|odds)"
    r".{0,16}(?:\b(?:0(?:\.\d+)?|1(?:\.0+)?)\b|\b\d+(?:\.\d+)?\s*[%％])"
    r"|(?:\b(?:0(?:\.\d+)?|1(?:\.0+)?)\b|\b\d+(?:\.\d+)?\s*[%％])"
    r".{0,16}(?:概率|可能性|置信度|probability|likelihood|confidence|odds)"
    r".{0,32}(?:厂家|制造商|源头|source\s*factory|manufacturer)"
)
FORBIDDEN_REPORT_KEY_RE = re.compile(
    r"(?i)\b(?:confidence_index|manufacturer_probability|source_factory_probability|"
    r"manufacturer_confidence|source_factory_confidence|factory_likelihood)\b"
)
FORBIDDEN_ATTRIBUTION_HEADER_RE = re.compile(
    r"(?i)(?:候选(?:厂家|制造商)|真实厂家|源头厂家|source\s*factory|manufacturer)"
    r".{0,80}(?:概率|可能性|厂家置信度|probability|likelihood|odds)"
    r"|(?:概率|可能性|厂家置信度|probability|likelihood|odds)"
    r".{0,80}(?:候选(?:厂家|制造商)|真实厂家|源头厂家|source\s*factory|manufacturer)"
)
FORBIDDEN_LEGACY_MULTI_AXIS_RE = re.compile(
    r"(?i)(?:"
    r"四(?:维|轴).{0,16}(?:概率|百分比|角色置信(?:度)?|置信度)"
    r"|综合(?:采购)?适配.{0,8}(?:估计|概率|置信(?:度)?)"
    r"|采购速查分.{0,80}(?:0\.\d+\s*[*×]?\s*[ABCD])"
    r"|four[-\s]?(?:axis|dimension(?:al)?).{0,24}(?:percent|probability|confidence)"
    r")"
)
FOUR_PERCENT_VALUES_RE = re.compile(
    r"(?:\b\d+(?:\.\d+)?\s*[%％].{0,32}){3}"
    r"\b\d+(?:\.\d+)?\s*[%％]"
)
PROBABILITY_NEGATION_RE = re.compile(
    r"(?i)(?:不输出|禁止|不提供|不采用|不使用|拒绝|已移除|已删除|已废弃|"
    r"未经(?:统计)?校准|not\s+(?:output|provide|use)|removed|deprecated)"
)


def _load_candidates(case_root: Path) -> list[dict[str, Any]]:
    candidate_root = case_root / "candidates"
    if not candidate_root.is_dir():
        raise FileNotFoundError(f"candidate directory is missing: {candidate_root}")
    candidates: list[dict[str, Any]] = []
    for path in sorted(candidate_root.glob("*.json")):
        value = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(value, dict):
            raise ValueError(f"candidate must be an object: {path}")
        value.setdefault("_source_path", path.relative_to(case_root).as_posix())
        candidates.append(value)
    return candidates


def _normalize_mainland(value: Any) -> str:
    if value is True or str(value).strip().lower() == "true":
        return "true"
    if value is False or str(value).strip().lower() == "false":
        return "false"
    return "unknown"


def _first_text(*values: Any) -> str:
    for value in values:
        if isinstance(value, str) and value.strip():
            return value.strip()
        if isinstance(value, list):
            rendered = "; ".join(str(item).strip() for item in value if str(item).strip())
            if rendered:
                return rendered
    return ""


def _contact_rows(case_root: Path) -> dict[str, list[dict[str, str]]]:
    path = case_root / "evidence" / "contacts.csv"
    if not path.is_file():
        return {}
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    grouped: dict[str, list[dict[str, str]]] = {}
    for row in rows:
        candidate_id = str(row.get("candidate_id", "")).strip()
        if candidate_id:
            raw_evidence_path = str(row.get("evidence_path", "")).strip()
            normalized_evidence_path = ""
            if raw_evidence_path:
                candidate_path = Path(raw_evidence_path)
                if not candidate_path.is_absolute() and ".." not in candidate_path.parts:
                    resolved = (case_root / candidate_path).resolve()
                    try:
                        resolved.relative_to(case_root)
                    except ValueError:
                        pass
                    else:
                        if resolved.is_file():
                            normalized_evidence_path = resolved.relative_to(case_root).as_posix()
            row["_safe_evidence_path"] = normalized_evidence_path
            grouped.setdefault(candidate_id, []).append(row)
    return grouped


def _contact_summary(
    rows: list[dict[str, str]],
) -> tuple[str, str, str, str, str, str]:
    values: list[str] = []
    source_urls: list[str] = []
    websites: list[str] = []
    verified: list[str] = []
    evidence_paths: list[str] = []
    statuses: list[str] = []
    for row in rows:
        status = str(row.get("status", "")).strip().casefold()
        source = sanitize_url_for_record(str(row.get("source_url", "")).strip())
        evidence_path = str(row.get("_safe_evidence_path", "")).strip()
        checked = str(row.get("last_verified_utc", "")).strip()
        if status not in {"verified", "current", "confirmed", "official"}:
            continue
        if not source or not evidence_path or not checked:
            continue
        statuses.append(status.upper())
        for label, field in (
            ("联系人", "contact_name"),
            ("电话", "phone"),
            ("邮箱", "email"),
            ("微信", "wechat"),
            ("WhatsApp", "whatsapp"),
            ("LinkedIn", "linkedin"),
        ):
            value = str(row.get(field, "")).strip()
            if value:
                values.append(f"{label}：{value}")
        source_urls.append(source)
        website = str(row.get("website", "")).strip()
        if website:
            websites.append(sanitize_url_for_record(website))
        evidence_paths.append(evidence_path)
        verified.append(checked)
    return (
        "；".join(dict.fromkeys(values)),
        "; ".join(dict.fromkeys(statuses)) if statuses else "UNVERIFIED_EXCLUDED" if rows else "",
        " ".join(dict.fromkeys(websites)),
        " ".join(dict.fromkeys(source_urls)),
        "; ".join(dict.fromkeys(evidence_paths)),
        max(verified) if verified else "",
    )


def _cert_state(result: dict[str, Any], scheme: str) -> str:
    state = result.get("certification_state", {})
    if isinstance(state, dict):
        value = state.get(scheme) or state.get(scheme.lower()) or state.get(scheme.upper())
        if isinstance(value, dict):
            return _first_text(value.get("state"), value.get("status"), "UNKNOWN")
        if value is not None:
            return str(value)
    return "UNKNOWN"


def _public_assessment(result: dict[str, Any]) -> dict[str, Any]:
    """Remove deprecated compatibility keys from public report payloads.

    中文：公开报告去除已弃用兼容字段，防止旧概率语义与新版多轴评估混用。
    """
    return {
        key: value
        for key, value in result.items()
        if key
        not in {
            "confidence_index",
            "manufacturer_probability",
            "source_factory_probability",
        }
    }


def _row_for_candidate(
    candidate: dict[str, Any], result: dict[str, Any], contacts: list[dict[str, str]]
) -> dict[str, Any]:
    (
        contact,
        contact_status,
        website,
        contact_source,
        contact_evidence_path,
        last_verified,
    ) = _contact_summary(contacts)
    evidence_urls = [
        sanitize_url_for_record(str(item.get("source_url", "")).strip())
        for item in candidate.get("evidence", [])
        if isinstance(item, dict) and str(item.get("source_url", "")).strip()
    ]
    critical_gap = _first_text(
        result.get("critical_gap"),
        result.get("next_required_evidence"),
        "; ".join(str(item.get("reason", "")) for item in result.get("caps", []) if item),
    )
    return {
        "candidate_id": candidate.get("candidate_id", "UNKNOWN"),
        "candidate": candidate.get("display_name", ""),
        "verified_mainland_site": _normalize_mainland(
            candidate.get(
                "verified_mainland_site",
                candidate.get("site_verification", {}).get("verified_mainland_site")
                if isinstance(candidate.get("site_verification"), dict)
                else None,
            )
        ),
        "target_role": _first_text(
            candidate.get("target_role"), candidate.get("roles"), candidate.get("role")
        ),
        "trace_stage": result.get("trace_stage", "S0_UNLINKED"),
        "analytic_confidence": result.get("analytic_confidence", "LOW"),
        "evidence_sufficiency_score": result.get(
            "evidence_sufficiency_score", result.get("raw_score", 0)
        ),
        "capability_fit_score": result.get("capability_fit_score", 0),
        "ul_state": _cert_state(result, "ul"),
        "watermark_state": _cert_state(result, "watermark"),
        "verification_priority": result.get("verification_priority", 0),
        "procurement_utility": result.get("procurement_utility", 0),
        "critical_gap": critical_gap,
        "contact": contact,
        "contact_status": contact_status,
        "website_or_product_page": website,
        "contact_source_url": contact_source,
        "contact_evidence_path": contact_evidence_path,
        "last_verified_utc": last_verified,
        "evidence_source_urls": " ".join(dict.fromkeys(evidence_urls)),
    }


def build_report_payload(case_root: Path) -> dict[str, Any]:
    case_root = case_root.resolve(strict=True)
    case_path = case_root / "case.json"
    case_record = (
        json.loads(case_path.read_text(encoding="utf-8")) if case_path.is_file() else {}
    )
    contacts_by_candidate = _contact_rows(case_root)
    rows: list[dict[str, Any]] = []
    assessments: list[dict[str, Any]] = []
    for candidate in _load_candidates(case_root):
        result = score_candidate(candidate)
        assessments.append(_public_assessment(result))
        rows.append(
            _row_for_candidate(
                candidate,
                result,
                contacts_by_candidate.get(str(candidate.get("candidate_id", "")), []),
            )
        )
    rows.sort(
        key=lambda item: (
            TRACE_STAGE_ORDER.get(str(item["trace_stage"]), -1),
            float(item["evidence_sufficiency_score"] or 0),
            float(item["capability_fit_score"] or 0),
        ),
        reverse=True,
    )
    mainland = [item for item in rows if item["verified_mainland_site"] == "true"]
    mainland.sort(
        key=lambda item: (
            float(item["procurement_utility"] or 0),
            float(item["verification_priority"] or 0),
            TRACE_STAGE_ORDER.get(str(item["trace_stage"]), -1),
        ),
        reverse=True,
    )
    return {
        "schema": 2,
        "report_semantics": "multi_axis_non_probability",
        "probability_notice": (
            "All numeric scores are heuristic evidence/capability/priority indices, "
            "not statistically calibrated manufacturer probabilities."
        ),
        "created_at_utc": utc_now(),
        "case_id": case_record.get("case_id", case_root.name),
        "case_root_label": case_root.name,
        "global_evidence_rows": rows,
        "mainland_procurement_rows": mainland,
        "assessments": assessments,
    }


def _md_cell(value: Any) -> str:
    rendered = str(value if value is not None else "")
    rendered = URL_RE.sub(lambda match: f"[{match.group(0)}]({match.group(0)})", rendered)
    return rendered.replace("|", "\\|").replace("\n", "<br>")


def _markdown_table(rows: list[dict[str, Any]]) -> str:
    columns = [(column, REPORT_LABELS[column]) for column in REPORT_COLUMNS]
    header = "| " + " | ".join(label for _, label in columns) + " |"
    divider = "|" + "|".join("---" for _ in columns) + "|"
    lines = [header, divider]
    for row in rows:
        lines.append("| " + " | ".join(_md_cell(row.get(key, "")) for key, _ in columns) + " |")
    if not rows:
        lines.append("| " + " | ".join("" for _ in columns) + " |")
    return "\n".join(lines)


def render_markdown(payload: dict[str, Any]) -> str:
    return (
        "# 源头制造商多轴核验报告\n\n"
        f"生成时间 UTC：`{payload['created_at_utc']}`\n\n"
        "> 本报告不输出未经统计校准的厂家概率。证据闭合度、能力适配和采购优先级相互独立。\n\n"
        "## 全球事实证据榜\n\n"
        + _markdown_table(payload["global_evidence_rows"])
        + "\n\n## 中国内地采购候选榜\n\n"
        + "仅包含 `verified_mainland_site=true` 的候选；地域不提高厂家归因。\n\n"
        + _markdown_table(payload["mainland_procurement_rows"])
        + "\n\n## 结论边界\n\n"
        "- `S0_UNLINKED`：未建立目标产品来源关系。\n"
        "- `S1_EXACT_PRODUCT_MATCH`：精确产品已匹配，但场址未闭合。\n"
        "- `S2_AUTHORIZED_SITE`：认证或审厂资料绑定准确场址。\n"
        "- `S3_PROCESS_CONFIRMED`：目标工序在准确场址获得确认。\n"
        "- `S4_BATCH_LINKED`：当前批次、工序、测试和出货链闭合。\n"
    )


def lint_report_text(text: str) -> list[str]:
    errors: list[str] = []
    if FORBIDDEN_PERCENT_RE.search(text):
        errors.append("un-calibrated manufacturer/source-factory probability value detected")
    if FORBIDDEN_REPORT_KEY_RE.search(text):
        errors.append("deprecated or probability-like report field detected")
    for line in text.splitlines():
        rendered = line.strip()
        if not rendered or PROBABILITY_NEGATION_RE.search(rendered):
            continue
        if FORBIDDEN_LEGACY_MULTI_AXIS_RE.search(rendered):
            errors.append("legacy multi-axis probability/confidence semantics detected")
            break
        if FOUR_PERCENT_VALUES_RE.search(rendered):
            errors.append("four-axis percentage row detected")
            break
        if FORBIDDEN_ATTRIBUTION_HEADER_RE.search(rendered):
            errors.append("manufacturer/source-factory probability column or heading detected")
            break
    return errors


def lint_report_payload(payload: dict[str, Any]) -> list[str]:
    errors = lint_report_text(json.dumps(payload, ensure_ascii=False))
    for section in ("global_evidence_rows", "mainland_procurement_rows"):
        for index, row in enumerate(payload.get(section, []), start=1):
            stage = str(row.get("trace_stage", ""))
            if stage not in TRACE_STAGE_ORDER:
                errors.append(f"{section}[{index}]: invalid trace_stage {stage!r}")
            if section == "mainland_procurement_rows" and row.get(
                "verified_mainland_site"
            ) != "true":
                errors.append(f"{section}[{index}]: mainland hard filter violated")
            contact = str(row.get("contact") or "").strip()
            if contact:
                if not str(row.get("contact_source_url") or "").strip():
                    errors.append(f"{section}[{index}]: contact has no source URL")
                if not str(row.get("contact_evidence_path") or "").strip():
                    errors.append(f"{section}[{index}]: contact has no frozen evidence path")
                else:
                    evidence_path = Path(str(row.get("contact_evidence_path") or ""))
                    if evidence_path.is_absolute() or ".." in evidence_path.parts:
                        errors.append(
                            f"{section}[{index}]: contact evidence path is not case-relative"
                        )
                if not str(row.get("last_verified_utc") or "").strip():
                    errors.append(f"{section}[{index}]: contact has no verification timestamp")
                statuses = {
                    item.strip().casefold()
                    for item in str(row.get("contact_status") or "").split(";")
                    if item.strip()
                }
                if not statuses or not statuses.issubset(
                    {"verified", "current", "confirmed", "official"}
                ):
                    errors.append(f"{section}[{index}]: contact status is not publishable")
    return errors


def lint_report_file(path: Path) -> list[str]:
    """Inspect every supported report format, including edited Office outputs.

    中文：逐格式检查报告，包括人工修改后的Office正文与表格；不能只验证最初的JSON数据。
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(value, dict) and "global_evidence_rows" in value:
            return lint_report_payload(value)
        return lint_report_text(json.dumps(value, ensure_ascii=False))
    if suffix in {".md", ".txt"}:
        return lint_report_text(path.read_text(encoding="utf-8-sig"))
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        payload = {
            "global_evidence_rows": [] if "mainland" in path.stem else rows,
            "mainland_procurement_rows": rows if "mainland" in path.stem else [],
        }
        return lint_report_payload(payload)
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:  # pragma: no cover - explicit CLI failure
            return [f"XLSX lint requires the reports extra: {error}"]
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sections: dict[str, list[dict[str, Any]]] = {
                "global_evidence_rows": [],
                "mainland_procurement_rows": [],
            }
            text_values: list[str] = []
            for worksheet in workbook.worksheets:
                values = list(worksheet.iter_rows(values_only=True))
                text_values.extend(
                    str(value) for row in values for value in row if value is not None
                )
                text_values.extend(
                    " | ".join(str(value or "") for value in row)
                    for row in values
                )
                if not values:
                    continue
                headers = [str(value or "") for value in values[0]]
                if not set(REPORT_COLUMNS).issubset(headers):
                    continue
                target = (
                    "mainland_procurement_rows"
                    if "内地" in worksheet.title or "mainland" in worksheet.title.casefold()
                    else "global_evidence_rows"
                )
                for values_row in values[1:]:
                    row = {
                        header: values_row[index] if index < len(values_row) else ""
                        for index, header in enumerate(headers)
                    }
                    if any(str(row.get(column, "") or "").strip() for column in REPORT_COLUMNS):
                        sections[target].append(row)
            return lint_report_text("\n".join(text_values)) + lint_report_payload(sections)
        finally:
            workbook.close()
    if suffix == ".docx":
        try:
            from docx import Document
            from docx.oxml.ns import qn
        except ImportError as error:  # pragma: no cover - explicit CLI failure
            return [f"DOCX lint requires the reports extra: {error}"]
        document = Document(path)
        text_nodes = [
            node.text or "" for node in document.element.iter(qn("w:t"))
        ]
        text_nodes.extend(
            " | ".join(cell.text.strip() for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        text_nodes.extend(
            str(relationship.target_ref)
            for relationship in document.part.rels.values()
            if relationship.reltype.endswith("/hyperlink")
        )
        inverse_labels = {value: key for key, value in REPORT_LABELS.items()}
        rows: list[dict[str, Any]] = []
        for table in document.tables:
            row: dict[str, Any] = {}
            for table_row in table.rows:
                if len(table_row.cells) != 2:
                    continue
                key = inverse_labels.get(table_row.cells[0].text.strip())
                if key:
                    row[key] = table_row.cells[1].text.strip()
            if row.get("candidate_id"):
                rows.append(row)
        return lint_report_text("\n".join(text_nodes)) + lint_report_payload(
            {"global_evidence_rows": rows, "mainland_procurement_rows": []}
        )
    return [
        f"unsupported report format {suffix or '<none>'}; use JSON, MD, TXT, CSV, XLSX, or DOCX"
    ]


def _write_xlsx(path: Path, payload: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as error:  # pragma: no cover - exercised by CLI failure path
        raise RuntimeError("XLSX output requires the 'reports' extra (openpyxl)") from error
    workbook = Workbook()
    default = workbook.active
    default.title = "全球事实证据榜"
    for worksheet, rows in (
        (default, payload["global_evidence_rows"]),
        (workbook.create_sheet("中国内地采购榜"), payload["mainland_procurement_rows"]),
    ):
        worksheet.append(REPORT_COLUMNS)
        for row in rows:
            worksheet.append([row.get(column, "") for column in REPORT_COLUMNS])
        for row_index in range(2, worksheet.max_row + 1):
            for column_index, column in enumerate(REPORT_COLUMNS, start=1):
                value = str(worksheet.cell(row_index, column_index).value or "")
                match = URL_RE.search(value)
                if match:
                    worksheet.cell(row_index, column_index).hyperlink = match.group(0)
                    worksheet.cell(row_index, column_index).style = "Hyperlink"
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = min(
                60, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
            )
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    links = workbook.create_sheet("链接明细")
    links.append(["candidate_id", "candidate", "field", "url"])
    seen_links: set[tuple[str, str, str]] = set()
    for row in payload["global_evidence_rows"]:
        for column in REPORT_COLUMNS:
            for match in URL_RE.finditer(str(row.get(column, "") or "")):
                key = (str(row.get("candidate_id", "")), column, match.group(0))
                if key in seen_links:
                    continue
                seen_links.add(key)
                links.append(
                    [row.get("candidate_id", ""), row.get("candidate", ""), column, match.group(0)]
                )
                cell = links.cell(links.max_row, 4)
                cell.hyperlink = match.group(0)
                cell.style = "Hyperlink"
    links.freeze_panes = "A2"
    links.auto_filter.ref = links.dimensions
    for cell in links[1]:
        cell.font = Font(bold=True)
    for letter, width in {"A": 18, "B": 30, "C": 28, "D": 80}.items():
        links.column_dimensions[letter].width = width
    for row in links.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def _write_docx(path: Path, payload: dict[str, Any]) -> None:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as error:  # pragma: no cover - exercised by CLI failure path
        raise RuntimeError("DOCX output requires the 'reports' extra (python-docx)") from error

    def set_font(run: Any, size: float = 10.5, bold: bool = False) -> None:
        run.font.name = "Calibri"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
        run.font.size = Pt(size)
        run.bold = bold

    def add_text_with_links(paragraph: Any, text: str, *, bold: bool = False) -> None:
        cursor = 0
        for match in URL_RE.finditer(text):
            if match.start() > cursor:
                set_font(paragraph.add_run(text[cursor : match.start()]), bold=bold)
            relationship_id = paragraph.part.relate_to(
                match.group(0),
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship_id)
            run = OxmlElement("w:r")
            properties = OxmlElement("w:rPr")
            style = OxmlElement("w:rStyle")
            style.set(qn("w:val"), "Hyperlink")
            properties.append(style)
            run.append(properties)
            text_element = OxmlElement("w:t")
            text_element.text = match.group(0)
            run.append(text_element)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
            cursor = match.end()
        if cursor < len(text):
            set_font(paragraph.add_run(text[cursor:]), bold=bold)

    def set_table_geometry(table: Any, widths: tuple[int, int] = (2700, 6660)) -> None:
        table.autofit = False
        table_pr = table._tbl.tblPr
        table_width = table_pr.first_child_found_in("w:tblW")
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_pr.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(sum(widths)))
        table_indent = table_pr.first_child_found_in("w:tblInd")
        if table_indent is None:
            table_indent = OxmlElement("w:tblInd")
            table_pr.append(table_indent)
        table_indent.set(qn("w:type"), "dxa")
        table_indent.set(qn("w:w"), "120")
        layout = table_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            table_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(width))
            grid.append(column)
        for row in table.rows:
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            for index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                cell_width = properties.first_child_found_in("w:tcW")
                if cell_width is None:
                    cell_width = OxmlElement("w:tcW")
                    properties.append(cell_width)
                cell_width.set(qn("w:type"), "dxa")
                cell_width.set(qn("w:w"), str(widths[index]))
                margins = properties.first_child_found_in("w:tcMar")
                if margins is None:
                    margins = OxmlElement("w:tcMar")
                    properties.append(margins)
                for side, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                    node = margins.find(qn(f"w:{side}"))
                    if node is None:
                        node = OxmlElement(f"w:{side}")
                        margins.append(node)
                    node.set(qn("w:w"), str(value))
                    node.set(qn("w:type"), "dxa")

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False
    document.settings.odd_and_even_pages_header_footer = True

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    title_style = document.styles["Title"]
    title_style.font.name = "Calibri"
    title_style.font.size = Pt(24)
    title_style.font.color.rgb = RGBColor.from_string("0B2545")
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(12)

    for header_container in (section.header, section.even_page_header):
        header = header_container.paragraphs[0]
        header.text = "Factory Trace Toolkit | 多轴证据报告"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            set_font(run, size=9)
            run.font.color.rgb = RGBColor.from_string("6B7280")
    for footer_container in (section.footer, section.even_page_footer):
        footer = footer_container.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer_run = footer.add_run("Factory Trace Toolkit 2.0")
        set_font(footer_run, size=9)
        footer_run.font.color.rgb = RGBColor.from_string("6B7280")

    document.add_heading("源头制造商多轴核验报告", level=0)
    metadata = document.add_paragraph()
    set_font(metadata.add_run(f"案件：{payload['case_id']}  |  生成时间 UTC：{payload['created_at_utc']}"), size=9.5)
    notice = document.add_paragraph(
        "本报告不输出未经统计校准的厂家概率。证据闭合度、能力适配和采购优先级相互独立。"
    )
    notice.style = document.styles["Normal"]
    for title, rows in (
        ("全球事实证据榜", payload["global_evidence_rows"]),
        ("中国内地采购候选榜", payload["mainland_procurement_rows"]),
    ):
        document.add_heading(title, level=1)
        if not rows:
            document.add_paragraph("无符合该榜单筛选条件的候选。")
        for row in rows:
            candidate_heading = document.add_heading(
                f"{row.get('candidate', '')} ({row.get('candidate_id', '')})", level=2
            )
            candidate_heading.paragraph_format.keep_with_next = True
            table = document.add_table(rows=len(REPORT_COLUMNS), cols=2)
            table.style = "Table Grid"
            set_table_geometry(table)
            for index, column in enumerate(REPORT_COLUMNS):
                label_cell, value_cell = table.rows[index].cells
                label_cell.text = ""
                label_paragraph = label_cell.paragraphs[0]
                set_font(label_paragraph.add_run(REPORT_LABELS[column]), bold=True)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8EEF5")
                label_cell._tc.get_or_add_tcPr().append(shading)
                value_cell.text = ""
                value_paragraph = value_cell.paragraphs[0]
                add_text_with_links(value_paragraph, str(row.get(column, "") or ""))
            document.add_paragraph()
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_report_formats(
    case_root: Path, output_dir: Path, formats: Iterable[str]
) -> dict[str, Any]:
    payload = build_report_payload(case_root)
    errors = lint_report_payload(payload)
    if errors:
        raise ValueError("report semantic lint failed: " + "; ".join(errors))
    output_dir.mkdir(parents=True, exist_ok=True)
    requested = {item.strip().lower() for item in formats if item.strip()}
    if not requested:
        requested = {"md", "json", "csv"}
    paths: list[Path] = []
    if "json" in requested:
        path = output_dir / "factory_trace_v2_report.json"
        atomic_write_json(path, payload)
        paths.append(path)
    if "md" in requested:
        path = output_dir / "factory_trace_v2_report.md"
        markdown = render_markdown(payload)
        text_errors = lint_report_text(markdown)
        if text_errors:
            raise ValueError("markdown semantic lint failed: " + "; ".join(text_errors))
        atomic_write_text(path, markdown)
        paths.append(path)
    if "csv" in requested:
        global_path = output_dir / "factory_trace_v2_global.csv"
        mainland_path = output_dir / "factory_trace_v2_mainland.csv"
        write_csv(global_path, REPORT_COLUMNS, payload["global_evidence_rows"])
        write_csv(mainland_path, REPORT_COLUMNS, payload["mainland_procurement_rows"])
        paths.extend([global_path, mainland_path])
    if "xlsx" in requested:
        path = output_dir / "factory_trace_v2_report.xlsx"
        _write_xlsx(path, payload)
        paths.append(path)
    if "docx" in requested:
        path = output_dir / "factory_trace_v2_report.docx"
        _write_docx(path, payload)
        paths.append(path)
    unsupported = sorted(requested - {"md", "json", "csv", "xlsx", "docx"})
    if unsupported:
        raise ValueError("unsupported report formats: " + ", ".join(unsupported))
    manifest = {
        "schema": 2,
        "created_at_utc": utc_now(),
        "report_semantics": payload["report_semantics"],
        "files": [
            {
                "path": path.relative_to(output_dir).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
            for path in paths
        ],
    }
    manifest_path = output_dir / "factory_trace_v2_report_manifest.json"
    atomic_write_json(manifest_path, manifest)
    manifest["manifest_path"] = str(manifest_path)
    return manifest


def run(args: object) -> int:
    formats: list[str] = []
    for value in getattr(args, "formats", []) or []:
        formats.extend(str(value).split(","))
    output_dir = getattr(args, "output_dir", None) or (
        Path(args.case_root) / "output" / "reports" / "v2"
    )
    result = write_report_formats(Path(args.case_root), Path(output_dir), formats)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def run_lint(args: object) -> int:
    path = Path(args.path)
    if path.suffix.lower() == ".json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        errors = lint_report_payload(payload)
    else:
        errors = lint_report_text(path.read_text(encoding="utf-8"))
    result = {"status": "PASS" if not errors else "FAIL", "path": str(path), "errors": errors}
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not errors else 1
