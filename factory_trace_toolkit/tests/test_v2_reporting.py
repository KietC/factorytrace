"""English: Regress report semantics, contact sourcing, links, and Office round trips.

中文：回归报告语义、联系方式来源、链接与 Office 回读；排序分数不是概率，手工编辑的旧概率列也必须拒绝。
"""
from __future__ import annotations

import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path


TOOLKIT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(TOOLKIT_ROOT / "src"))

from factorytrace.init_case import create_case  # noqa: E402
from factorytrace.reporting import (  # noqa: E402
    build_report_payload,
    lint_report_file,
    lint_report_payload,
    lint_report_text,
    write_report_formats,
)


class V2ReportingTests(unittest.TestCase):
    def test_report_has_multi_axis_semantics_and_mainland_hard_filter(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            case = create_case("report", root)
            first_path = case / "candidates" / "CAND-001.json"
            first = json.loads(first_path.read_text(encoding="utf-8"))
            first["display_name"] = "Mainland candidate"
            first["verified_mainland_site"] = True
            first_path.write_text(json.dumps(first, ensure_ascii=False), encoding="utf-8")
            second = dict(first)
            second["candidate_id"] = "CAND-002"
            second["display_name"] = "Unknown-site candidate"
            second["verified_mainland_site"] = "unknown"
            (case / "candidates" / "CAND-002.json").write_text(
                json.dumps(second, ensure_ascii=False), encoding="utf-8"
            )
            payload = build_report_payload(case)
            self.assertEqual(payload["report_semantics"], "multi_axis_non_probability")
            self.assertEqual(len(payload["global_evidence_rows"]), 2)
            self.assertEqual(len(payload["mainland_procurement_rows"]), 1)
            self.assertEqual(
                payload["mainland_procurement_rows"][0]["candidate_id"], "CAND-001"
            )
            rendered = json.dumps(payload, ensure_ascii=False)
            self.assertNotIn("manufacturer_probability", rendered)
            self.assertNotIn(str(case.resolve()), rendered)

    def test_report_lint_rejects_probability_and_unsourced_contact(self) -> None:
        self.assertTrue(lint_report_text("示例企业真实厂家概率 59%"))
        self.assertTrue(lint_report_text("示例企业真实厂家置信度 59％"))
        self.assertTrue(lint_report_text("源头厂家可能性 0.59"))
        self.assertTrue(lint_report_text("source factory confidence: 0.59"))
        self.assertTrue(lint_report_text("| 候选厂家 | 概率 | 联系方式 |"))
        self.assertFalse(lint_report_text("本报告不输出未经统计校准的厂家概率。"))
        self.assertFalse(lint_report_text("分析置信度：LOW；ESS：59"))
        payload = {
            "global_evidence_rows": [
                {
                    "trace_stage": "S0_UNLINKED",
                    "verified_mainland_site": "true",
                    "contact": "+86 123",
                    "contact_source_url": "",
                    "contact_status": "unverified",
                    "contact_evidence_path": "",
                    "last_verified_utc": "",
                }
            ],
            "mainland_procurement_rows": [],
        }
        self.assertTrue(
            any("contact has no source" in item for item in lint_report_payload(payload))
        )

    def test_write_base_formats_and_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            case = create_case("formats", root)
            output = root / "out"
            manifest = write_report_formats(case, output, ["md", "json", "csv"])
            self.assertEqual(len(manifest["files"]), 4)
            for record in manifest["files"]:
                self.assertEqual(len(record["sha256"]), 64)
            with (output / "factory_trace_v2_global.csv").open(
                encoding="utf-8-sig", newline=""
            ) as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual(len(rows), 1)

    def test_write_xlsx_and_docx_from_same_payload(self) -> None:
        try:
            import openpyxl  # noqa: F401
            import docx  # noqa: F401
        except ImportError:
            self.skipTest("reports extra is not installed")
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            case = create_case("office-formats", root)
            contacts = case / "evidence" / "contacts.csv"
            evidence_file = case / "evidence" / "communications" / "contact.pdf"
            evidence_file.write_bytes(b"synthetic frozen contact evidence")
            contacts.write_text(
                "candidate_id,contact_name,phone,email,wechat,whatsapp,linkedin,website,source_url,evidence_path,last_verified_utc,status,notes\n"
                "CAND-001,Alice,+86 123,a@example.test,,,https://linkedin.example/alice,"
                "https://factory.example,https://registry.example/contact?token=SECRET&view=public,evidence/communications/contact.pdf,2026-08-01T00:00:00Z,verified,test\n",
                encoding="utf-8-sig",
            )
            output = root / "out"
            manifest = write_report_formats(case, output, ["md", "xlsx", "docx"])
            self.assertEqual(len(manifest["files"]), 3)
            self.assertTrue((output / "factory_trace_v2_report.xlsx").is_file())
            self.assertTrue((output / "factory_trace_v2_report.docx").is_file())
            markdown = (output / "factory_trace_v2_report.md").read_text(encoding="utf-8")
            self.assertIn("[https://factory.example](https://factory.example)", markdown)
            from docx import Document
            from openpyxl import load_workbook

            document = Document(output / "factory_trace_v2_report.docx")
            document_text = "\n".join(
                paragraph.text for table in document.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs
            )
            self.assertIn("最后核验时间 UTC", document_text)
            self.assertIn("https://registry.example/contact", document_text)
            self.assertNotIn("SECRET", document_text)
            hyperlink_relationships = [
                relationship
                for relationship in document.part.rels.values()
                if relationship.reltype.endswith("/hyperlink")
            ]
            self.assertTrue(hyperlink_relationships)
            workbook = load_workbook(output / "factory_trace_v2_report.xlsx", read_only=False)
            self.assertIn("链接明细", workbook.sheetnames)
            self.assertGreaterEqual(workbook["链接明细"].max_row, 4)
            self.assertFalse(lint_report_file(output / "factory_trace_v2_report.xlsx"))
            self.assertFalse(lint_report_file(output / "factory_trace_v2_report.docx"))
            workbook["全球事实证据榜"]["B2"] = "源头厂家置信度 59％"
            workbook.save(output / "factory_trace_v2_report.xlsx")
            workbook.close()
            self.assertTrue(lint_report_file(output / "factory_trace_v2_report.xlsx"))
            document.add_paragraph("source factory confidence: 0.59")
            document.save(output / "factory_trace_v2_report.docx")
            self.assertTrue(lint_report_file(output / "factory_trace_v2_report.docx"))

    def test_legacy_probability_columns_are_rejected_in_office_tables(self) -> None:
        try:
            from docx import Document
            from openpyxl import Workbook
        except ImportError:
            self.skipTest("reports extra is not installed")
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["候选厂家", "概率", "网址"])
            sheet.append(["示例候选", "68%", "https://example.test"])
            xlsx = root / "legacy.xlsx"
            workbook.save(xlsx)
            workbook.close()
            self.assertTrue(lint_report_file(xlsx))

            document = Document()
            table = document.add_table(rows=2, cols=3)
            for index, value in enumerate(["候选厂家", "概率", "网址"]):
                table.rows[0].cells[index].text = value
            for index, value in enumerate(["示例候选", "68%", "https://example.test"]):
                table.rows[1].cells[index].text = value
            docx = root / "legacy.docx"
            document.save(docx)
            self.assertTrue(lint_report_file(docx))


if __name__ == "__main__":
    unittest.main()
