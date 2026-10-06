#!/usr/bin/env python3
"""English: Probe real OCR, media, schema, process-pack, and Office capabilities.

中文：以临时合成材料验证 OCR、媒体、Schema、工艺包与 Office 能力；包能导入不等于功能可用，输出回执可能含环境路径。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import time
import tomllib
import traceback
from importlib import metadata
from pathlib import Path
from typing import Any, Callable

from factorytrace.common import (
    atomic_write_json,
    configure_utf8_stdio,
    sha256_file,
    utc_now,
)
from factorytrace.environment import _tool


def package_version(name: str) -> str | None:
    try:
        return metadata.version(name)
    except metadata.PackageNotFoundError:
        return None


def executable_record(path: str | None) -> dict[str, Any]:
    if not path:
        return {"found": False, "path": None, "sha256": None}
    resolved = Path(path).resolve()
    return {
        "found": resolved.is_file(),
        "path": str(resolved),
        "sha256": sha256_file(resolved) if resolved.is_file() else None,
    }


def run_command(command: list[str], *, timeout: int = 30) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        command,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=timeout,
    )


def check_pip() -> dict[str, Any]:
    result = run_command([sys.executable, "-m", "pip", "check"], timeout=60)
    return {"stdout": result.stdout.strip(), "returncode": result.returncode}


def check_tested_lock(work_root: Path) -> dict[str, Any]:
    constraints = work_root / "constraints-tested.txt"
    expected: dict[str, str] = {}
    for raw_line in constraints.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if "==" not in line:
            raise ValueError(f"constraint is not exact: {line}")
        name, version = line.split("==", 1)
        expected[name.strip()] = version.strip()
    observed = {name: package_version(name) for name in expected}
    mismatches = {
        name: {"expected": expected[name], "observed": observed[name]}
        for name in expected
        if observed[name] != expected[name]
    }
    if mismatches:
        raise RuntimeError(f"tested lock mismatch: {mismatches}")
    pyproject = tomllib.loads((work_root / "pyproject.toml").read_text(encoding="utf-8"))
    project = pyproject["project"]
    installed_metadata = metadata.metadata("factory-trace-toolkit")
    expected_python = project["requires-python"]
    observed_python = installed_metadata.get("Requires-Python")
    normalize_specifiers = lambda value: tuple(  # noqa: E731
        sorted(item.strip() for item in str(value or "").split(",") if item.strip())
    )
    if normalize_specifiers(observed_python) != normalize_specifiers(expected_python):
        raise RuntimeError(
            "installed toolkit Requires-Python differs from source: "
            f"{observed_python!r} != {expected_python!r}"
        )

    def normalized_requirement(value: str) -> tuple[str, tuple[str, ...]]:
        requirement = value.split(";", 1)[0].strip()
        name = requirement
        for delimiter in ("<", ">", "=", "!", "~", "[", " "):
            name = name.split(delimiter, 1)[0]
        specifier_text = requirement[len(name) :].strip()
        specifiers = tuple(
            sorted(item.strip().casefold() for item in specifier_text.split(",") if item.strip())
        )
        return name.casefold().replace("_", "-"), specifiers

    expected_requirements = {
        normalized_requirement(item)
        for item in (
            list(project.get("dependencies", []))
            + list(project.get("optional-dependencies", {}).get("full", []))
        )
    }
    installed_requirements = {
        normalized_requirement(item)
        for item in (installed_metadata.get_all("Requires-Dist") or [])
    }
    missing_metadata = expected_requirements - installed_requirements
    if missing_metadata:
        raise RuntimeError(
            f"installed toolkit Requires-Dist metadata is stale/incomplete: {sorted(missing_metadata)}"
        )
    return {
        "constraints": str(constraints.resolve()),
        "packages_checked": len(expected),
        "versions": observed,
        "toolkit_metadata": {
            "requires_python": observed_python,
            "required_distributions": [
                {"name": name, "specifiers": list(specifiers)}
                for name, specifiers in sorted(expected_requirements)
            ],
        },
    }


def check_jsonschema(work_root: Path) -> dict[str, Any]:
    from jsonschema import Draft202012Validator, validate

    schema = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "type": "object",
        "required": ["trace_stage"],
        "properties": {
            "trace_stage": {
                "enum": [
                    "S0_UNLINKED",
                    "S1_EXACT_PRODUCT_MATCH",
                    "S2_AUTHORIZED_SITE",
                    "S3_PROCESS_CONFIRMED",
                    "S4_BATCH_LINKED",
                ]
            }
        },
        "additionalProperties": False,
    }
    Draft202012Validator.check_schema(schema)
    validate({"trace_stage": "S0_UNLINKED"}, schema)
    schema_files = sorted((work_root / "schemas").glob("*.schema.json"))
    for path in schema_files:
        Draft202012Validator.check_schema(json.loads(path.read_text(encoding="utf-8")))
    return {"draft": "2020-12", "workspace_schemas_checked": len(schema_files)}


def check_pillow_and_opencv(temp: Path) -> dict[str, Any]:
    import cv2
    import numpy as np
    from PIL import Image

    image = np.full((300, 1200, 3), 255, dtype=np.uint8)
    cv2.putText(
        image,
        "FACTORY 316",
        (35, 195),
        cv2.FONT_HERSHEY_SIMPLEX,
        3.4,
        (0, 0, 0),
        8,
        cv2.LINE_AA,
    )
    png = temp / "factory-316.png"
    jpeg = temp / "factory-316.jpg"
    if not cv2.imwrite(str(png), image):
        raise RuntimeError("OpenCV failed to write the PNG probe")
    loaded = cv2.imread(str(png), cv2.IMREAD_COLOR)
    if loaded is None or loaded.shape != image.shape:
        raise RuntimeError("OpenCV failed to decode the generated PNG")
    edges = cv2.Canny(loaded, 80, 160)
    if int(np.count_nonzero(edges)) < 100:
        raise RuntimeError("OpenCV edge detector returned an implausibly empty result")
    with Image.open(png) as source:
        source.convert("RGB").save(jpeg, quality=95)
    with Image.open(jpeg) as reopened:
        reopened.load()
        dimensions = [reopened.width, reopened.height]
    return {
        "png": str(png),
        "jpeg": str(jpeg),
        "dimensions": dimensions,
        "edge_nonzero": int(np.count_nonzero(edges)),
        "pixel_sha256": hashlib.sha256(loaded.tobytes()).hexdigest(),
    }


def check_tesseract(temp: Path, work_root: Path) -> dict[str, Any]:
    executable = _tool("tesseract")
    if not executable:
        raise FileNotFoundError("tesseract executable was not found")
    version = run_command([executable, "--version"]).stdout.splitlines()[0]
    source_tessdata = work_root / "resources" / "tessdata"
    manifest = json.loads(
        (source_tessdata / "MANIFEST.json").read_text(encoding="utf-8")
    )
    tessdata = temp / "tessdata"
    tessdata.mkdir()
    verified_hashes: dict[str, str] = {}
    for name, expected in manifest["files"].items():
        source = source_tessdata / name
        observed = sha256_file(source).upper()
        if observed != str(expected).upper():
            raise RuntimeError(f"tessdata hash mismatch for {name}: {observed}")
        verified_hashes[name] = observed
        if source.suffix == ".traineddata":
            shutil.copy2(source, tessdata / name)
    listed = run_command(
        [executable, "--tessdata-dir", str(tessdata), "--list-langs"]
    )
    languages = sorted(
        line.strip()
        for line in listed.stdout.splitlines()
        if line.strip() and not line.lower().startswith("list of available languages")
    )
    required_languages = {"eng", "osd", "chi_sim", "chi_tra"}
    if not required_languages.issubset(languages):
        raise RuntimeError(f"tessdata languages missing: {required_languages - set(languages)}")
    result = run_command(
        [
            executable,
            str(temp / "factory-316.png"),
            "stdout",
            "--tessdata-dir",
            str(tessdata),
            "--psm",
            "7",
            "-l",
            "eng",
        ]
    )
    text = " ".join(result.stdout.upper().split())
    if "FACTORY" not in text or "316" not in text:
        raise RuntimeError(f"OCR mismatch: {text!r}")
    chinese_results: dict[str, str] = {}
    probe_root = work_root / "resources" / "ocr_probes"
    probe_manifest = json.loads(
        (probe_root / "MANIFEST.json").read_text(encoding="utf-8")
    )
    for language, probe, expected in (
        ("chi_sim", "chi_sim_probe.png", "工厂生产316不锈钢"),
        ("chi_tra", "chi_tra_probe.png", "工廠生產316不鏽鋼"),
    ):
        expected_probe_hash = probe_manifest["files"][probe]["sha256"]
        observed_probe_hash = sha256_file(probe_root / probe).upper()
        if observed_probe_hash != expected_probe_hash.upper():
            raise RuntimeError(f"OCR probe hash mismatch for {probe}: {observed_probe_hash}")
        output = run_command(
            [
                executable,
                str(probe_root / probe),
                "stdout",
                "--tessdata-dir",
                str(tessdata),
                "--psm",
                "7",
                "-l",
                language,
            ]
        ).stdout
        normalized = "".join(output.split())
        if expected not in normalized:
            raise RuntimeError(f"{language} OCR mismatch: {normalized!r}")
        chinese_results[language] = normalized
    return {
        "version": version,
        "languages": languages,
        "recognized_text": {"eng": text, **chinese_results},
        "tessdata_revision": manifest["revision"],
        "tessdata_hashes": verified_hashes,
        **executable_record(executable),
    }


def check_exiftool(temp: Path) -> dict[str, Any]:
    executable = _tool("exiftool")
    if not executable:
        raise FileNotFoundError("exiftool executable was not found")
    version = run_command([executable, "-ver"]).stdout.strip()
    result = run_command(
        [executable, "-j", "-MIMEType", "-ImageWidth", "-ImageHeight", str(temp / "factory-316.jpg")]
    )
    values = json.loads(result.stdout)
    if not values or values[0].get("MIMEType") != "image/jpeg":
        raise RuntimeError(f"unexpected ExifTool output: {values!r}")
    if [values[0].get("ImageWidth"), values[0].get("ImageHeight")] != [1200, 300]:
        raise RuntimeError(f"unexpected ExifTool dimensions: {values!r}")
    return {"version": version, "metadata": values[0], **executable_record(executable)}


def check_ffmpeg(temp: Path) -> dict[str, Any]:
    ffmpeg = _tool("ffmpeg")
    ffprobe = _tool("ffprobe")
    if not ffmpeg or not ffprobe:
        raise FileNotFoundError("ffmpeg and ffprobe are both required")
    video = temp / "probe.mkv"
    run_command(
        [
            ffmpeg,
            "-hide_banner",
            "-loglevel",
            "error",
            "-f",
            "lavfi",
            "-i",
            "color=c=blue:s=320x240:d=1",
            "-c:v",
            "ffv1",
            "-y",
            str(video),
        ],
        timeout=60,
    )
    probe = run_command(
        [
            ffprobe,
            "-v",
            "error",
            "-select_streams",
            "v:0",
            "-show_entries",
            "stream=codec_name,width,height,duration",
            "-of",
            "json",
            str(video),
        ]
    )
    values = json.loads(probe.stdout)
    stream = values["streams"][0]
    if stream.get("width") != 320 or stream.get("height") != 240:
        raise RuntimeError(f"unexpected ffprobe dimensions: {stream!r}")
    version = run_command([ffmpeg, "-version"]).stdout.splitlines()[0]
    return {
        "version": version,
        "stream": stream,
        "video_bytes": video.stat().st_size,
        "ffmpeg": executable_record(ffmpeg),
        "ffprobe": executable_record(ffprobe),
    }


def check_reports(temp: Path) -> dict[str, Any]:
    from docx import Document
    from openpyxl import Workbook, load_workbook

    xlsx = temp / "report-smoke.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Evidence"
    worksheet.append(["candidate", "trace_stage"])
    worksheet.append(["FACTORY 316", "S0_UNLINKED"])
    workbook.save(xlsx)
    reopened = load_workbook(xlsx, read_only=True, data_only=True)
    values = list(reopened["Evidence"].values)
    reopened.close()
    if values[1] != ("FACTORY 316", "S0_UNLINKED"):
        raise RuntimeError(f"XLSX round trip mismatch: {values!r}")

    docx = temp / "report-smoke.docx"
    document = Document()
    document.add_heading("Factory Trace", level=1)
    document.add_paragraph("FACTORY 316 / S0_UNLINKED")
    document.save(docx)
    reopened_document = Document(docx)
    text = "\n".join(paragraph.text for paragraph in reopened_document.paragraphs)
    if "FACTORY 316" not in text or "S0_UNLINKED" not in text:
        raise RuntimeError(f"DOCX round trip mismatch: {text!r}")
    return {
        "xlsx_round_trip": True,
        "docx_round_trip": True,
        "xlsx_bytes": xlsx.stat().st_size,
        "docx_bytes": docx.stat().st_size,
    }


def check_yaml_process_pack(temp: Path) -> dict[str, Any]:
    from factorytrace.process_packs import load_process_pack

    path = temp / "yaml_probe.yaml"
    path.write_text(
        """schema: factorytrace.process-pack.v2
pack_id: yaml_probe
version: 1
description: YAML parser capability probe
required_processes: [forming]
required_evidence: [work_order]
required_photos: [machine_nameplate]
required_documents: [inspection_record]
certification: {mode: all, schemes: [], hard_gates: []}
events: {required_biz_steps: [forming], certified_biz_steps: []}
social: {maximum_authority: entity_identity}
""",
        encoding="utf-8",
    )
    pack = load_process_pack("yaml_probe", pack_dir=temp)
    if pack.get("pack_id") != "yaml_probe" or pack.get("required_processes") != ["forming"]:
        raise RuntimeError(f"YAML process-pack round trip mismatch: {pack!r}")
    return {"pack_id": pack["pack_id"], "source_path": pack["source_path"]}


def execute_check(name: str, function: Callable[[], dict[str, Any]]) -> dict[str, Any]:
    started = time.perf_counter()
    try:
        details = function()
        return {
            "status": "PASS",
            "duration_seconds": round(time.perf_counter() - started, 3),
            "details": details,
        }
    except Exception as error:  # Report all dependency failures. / 中文：一次记录全部依赖失败，便于集中修复。
        return {
            "status": "FAIL",
            "duration_seconds": round(time.perf_counter() - started, 3),
            "error": f"{type(error).__name__}: {error}",
            "traceback": traceback.format_exc(),
        }


def main() -> int:
    configure_utf8_stdio()
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--work-root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(prefix="factorytrace-capability-") as temporary:
        temp = Path(temporary)
        checks: dict[str, dict[str, Any]] = {}
        checks["pip_integrity"] = execute_check("pip_integrity", check_pip)
        checks["tested_lock"] = execute_check(
            "tested_lock", lambda: check_tested_lock(args.work_root.resolve())
        )
        checks["jsonschema"] = execute_check(
            "jsonschema", lambda: check_jsonschema(args.work_root.resolve())
        )
        checks["pillow_opencv"] = execute_check(
            "pillow_opencv", lambda: check_pillow_and_opencv(temp)
        )
        checks["tesseract_ocr"] = execute_check(
            "tesseract_ocr", lambda: check_tesseract(temp, args.work_root.resolve())
        )
        checks["exiftool_metadata"] = execute_check(
            "exiftool_metadata", lambda: check_exiftool(temp)
        )
        checks["ffmpeg_decode_probe"] = execute_check(
            "ffmpeg_decode_probe", lambda: check_ffmpeg(temp)
        )
        checks["report_round_trip"] = execute_check(
            "report_round_trip", lambda: check_reports(temp)
        )
        checks["yaml_process_pack"] = execute_check(
            "yaml_process_pack", lambda: check_yaml_process_pack(temp)
        )

    failures = [name for name, value in checks.items() if value["status"] != "PASS"]
    report = {
        "schema": 1,
        "captured_at_utc": utc_now(),
        "status": "PASS" if not failures else "FAIL",
        "python": {
            "version": sys.version,
            "executable": sys.executable,
        },
        "packages": {
            name: package_version(name)
            for name in (
                "factory-trace-toolkit",
                "Pillow",
                "jsonschema",
                "numpy",
                "opencv-python-headless",
                "openpyxl",
                "python-docx",
                "PyYAML",
                "pip",
                "setuptools",
            )
        },
        "checks": checks,
        "failures": failures,
        "local_model_required": False,
        "local_model_note": (
            "Source-factory inference is evidence/rules based. A local model or GPU is "
            "optional enrichment and is not a runtime dependency."
        ),
    }
    atomic_write_json(args.output, report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())
