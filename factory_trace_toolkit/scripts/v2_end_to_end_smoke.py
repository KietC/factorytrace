#!/usr/bin/env python3
"""English: Exercise the versioned CLI end to end in a disposable synthetic case.

中文：在临时合成案件中端到端执行新版 CLI；记录退出码与生成结果，不读取真实供应商或客户案件。
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

from factorytrace.common import atomic_write_json, configure_utf8_stdio, utc_now


def invoke(arguments: list[str], records: list[dict[str, Any]]) -> dict[str, Any]:
    started = time.perf_counter()
    command = [sys.executable, "-m", "factorytrace", *arguments]
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    environment["PYTHONIOENCODING"] = "utf-8"
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        env=environment,
        timeout=90,
    )
    record = {
        "arguments": arguments,
        "returncode": result.returncode,
        "duration_seconds": round(time.perf_counter() - started, 3),
        "stdout": result.stdout.strip(),
        "stderr": result.stderr.strip(),
    }
    records.append(record)
    if result.returncode != 0:
        raise RuntimeError(
            f"command failed ({result.returncode}): {arguments!r}\n{result.stderr}\n{result.stdout}"
        )
    return json.loads(result.stdout) if result.stdout.strip().startswith(("{", "[")) else record


def build_candidate(case_root: Path) -> None:
    path = case_root / "candidates" / "CAND-001.json"
    candidate = json.loads(path.read_text(encoding="utf-8"))
    candidate.update(
        {
            "display_name": "Synthetic Capable Forming Site",
            "target_processes": ["forming"],
            "verified_mainland_site": True,
            "parties": [
                {
                    "party_id": "PARTY-SMOKE-001",
                    "legal_name": "Synthetic Capable Forming Site",
                    "roles": ["claimed_manufacturer"],
                }
            ],
            "sites": [
                {
                    "site_id": "SITE-SMOKE-001",
                    "address": "Jiangmen, Guangdong, China",
                    "verification_state": "verified",
                }
            ],
            "process_scope": [
                {
                    "process": "forming",
                    "performance_state": "observed",
                    "make_or_buy": "unknown",
                }
            ],
            "evidence": [
                {
                    "schema_version": 2,
                    "evidence_id": "EV-SMOKE-001",
                    "dimension": "physical_site_process",
                    "stance": "support",
                    "strength": "direct",
                    "source_class": "government",
                    "independence_group": "GOV-SMOKE-001",
                    "entity_bind": "exact",
                    "site_bind": "exact",
                    "sku_bind": "none",
                    "process": "forming",
                    "source_url": "https://example.invalid/government/smoke",
                    "captured_at_utc": "2026-08-01T00:00:00Z",
                    "observed_fact": "The site has metal-forming equipment.",
                    "limitations": "No target product, SKU, order, tooling or batch link.",
                    "review_status": "reviewed",
                }
            ],
        }
    )
    candidate["entities"].update(
        {
            "manufacturer_entity": "Synthetic Capable Forming Site",
            "manufacturing_sites": ["SITE-SMOKE-001"],
        }
    )
    path.write_text(json.dumps(candidate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    contacts = case_root / "evidence" / "contacts.csv"
    with contacts.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=[
                "candidate_id",
                "role",
                "contact_name",
                "phone",
                "email",
                "wechat",
                "whatsapp",
                "linkedin",
                "website",
                "source_url",
                "evidence_path",
                "last_verified_utc",
                "status",
                "notes",
            ],
        )
        writer.writeheader()
        writer.writerow(
            {
                "candidate_id": "CAND-001",
                "role": "public office",
                "phone": "+86 750 000 0000",
                "website": "https://example.invalid/factory",
                "source_url": "https://example.invalid/contact",
                "last_verified_utc": "2026-08-01T00:00:00Z",
                "status": "unverified",
            }
        )


def main() -> int:
    configure_utf8_stdio()
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    records: list[dict[str, Any]] = []
    assertions: list[str] = []
    with tempfile.TemporaryDirectory(prefix="factorytrace-v2-e2e-") as temporary:
        root = Path(temporary)
        invoke(["init", "smoke-case", "--root", str(root)], records)
        case = root / "smoke-case"
        build_candidate(case)
        receipt = case / "work" / "blocked.json"
        receipt.write_text(
            json.dumps(
                {
                    "receipt_id": "ACCESS-SMOKE-001",
                    "url": "https://example.invalid/login",
                    "access_status": "blocked_login",
                }
            ),
            encoding="utf-8",
        )

        validation = invoke(["validate", "--case-root", str(case)], records)
        assert validation["status"] == "PASS"
        assertions.append("schema_validation_pass")

        assessment_index = invoke(["assess", "--case-root", str(case)], records)
        candidate = assessment_index["candidates"][0]
        assert candidate["trace_stage"] == "S0_UNLINKED"
        assert candidate["capability_fit_score"] >= 90
        assertions.append("capable_site_remains_unlinked")

        hypotheses = invoke(["hypotheses", "--case-root", str(case)], records)
        assert hypotheses["candidates"][0]["ach"]["leading_hypotheses"] == ["H2"]
        assertions.append("ach_selects_capability_without_supply_link")

        invoke(
            ["queries", "--case-root", str(case), "--pack", "metal_forming"],
            records,
        )
        query_doc = json.loads(
            (case / "work" / "queries" / "TARGET-001_queries.json").read_text(
                encoding="utf-8"
            )
        )
        assert query_doc["process_packs"] == ["metal_forming"]
        assertions.append("process_pack_queries_generated")

        invoke(
            ["capture", "--case-root", str(case), "--receipt", str(receipt)], records
        )
        access = json.loads(
            (case / "evidence" / "access_receipts.jsonl").read_text(encoding="utf-8").splitlines()[0]
        )
        assert access["evidence_eligible"] is False
        assert access["absence_conclusion_permitted"] is False
        assertions.append("blocked_access_is_not_absence")

        invoke(["social", "assess", "--case-root", str(case)], records)
        invoke(["events", "validate", "--case-root", str(case)], records)
        invoke(["process-pack", "list"], records)
        invoke(
            [
                "report",
                "--case-root",
                str(case),
                "--format",
                "md,json,csv,xlsx,docx",
            ],
            records,
        )
        invoke(["lint-report", "--case-root", str(case)], records)
        invoke(["audit", "--case-root", str(case), "--stage", "research"], records)

        report_dir = case / "output" / "reports" / "v2"
        from docx import Document
        from openpyxl import load_workbook

        workbook = load_workbook(
            report_dir / "factory_trace_v2_report.xlsx", read_only=True, data_only=True
        )
        assert workbook["全球事实证据榜"]["A2"].value == "CAND-001"
        workbook.close()
        document = Document(report_dir / "factory_trace_v2_report.docx")
        assert "源头制造商多轴核验报告" in "\n".join(
            paragraph.text for paragraph in document.paragraphs
        )
        report_text = (report_dir / "factory_trace_v2_report.md").read_text(encoding="utf-8")
        assert "真实厂家概率" not in report_text
        report_payload = json.loads(
            (report_dir / "factory_trace_v2_report.json").read_text(encoding="utf-8")
        )
        assert report_payload["global_evidence_rows"][0]["contact"] == ""
        assert report_payload["global_evidence_rows"][0]["contact_status"] == "UNVERIFIED_EXCLUDED"
        assertions.append("all_report_formats_round_trip_and_lint")

    report = {
        "schema": 2,
        "status": "PASS",
        "captured_at_utc": utc_now(),
        "commands_run": len(records),
        "assertions": assertions,
        "commands": records,
        "temporary_case_removed": True,
    }
    atomic_write_json(args.output, report)
    print(
        json.dumps(
            {key: value for key, value in report.items() if key != "commands"},
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
